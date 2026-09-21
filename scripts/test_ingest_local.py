# -*- coding: utf-8 -*-
"""知识库上传（ingest）路径实测：md / txt / docx / 不支持的扩展名，以及来源标注。"""
import io
import json
import os
import time

import httpx

BASE = "http://127.0.0.1:8000"
WORK = r"D:\dsh\ai-robot-agent\.logs\ktest"
os.makedirs(WORK, exist_ok=True)

MD_PATH = os.path.join(WORK, "会员积分规则.md")
TXT_PATH = os.path.join(WORK, "发票与开票.txt")
DOCX_PATH = os.path.join(WORK, "配送范围说明.docx")
BAD_PATH = os.path.join(WORK, "不该支持.csv")

with io.open(MD_PATH, "w", encoding="utf-8") as f:
    f.write("""# 会员积分规则

## 积分怎么获得
1. 每完成一笔订单，按实付金额 1 元累积 1 积分；
2. 每日签到可额外获得 5 积分，连续签到 7 天再奖励 50 积分。

## 积分怎么使用
1. 100 积分可抵扣 1 元，单笔订单最多抵扣订单金额的 20%；
2. 积分不可提现、不可转让，有效期 12 个月。

## 积分什么时候到账
订单确认收货后 24 小时内到账；若发生退款，对应积分同步扣回。
""")

with io.open(TXT_PATH, "w", encoding="utf-8") as f:
    f.write("""发票与开票说明

本平台支持开具电子普通发票，下单时在结算页勾选"需要发票"并填写抬头与税号。
发票金额为实付金额，不含平台优惠券抵扣部分。
开票申请提交后 3 个工作日内发送到预留邮箱；已开票订单如需重开，请联系客服并提供订单号。
""")

from docx import Document  # noqa: E402

doc = Document()
doc.add_paragraph("配送范围说明")
doc.add_paragraph("平台默认使用顺丰或中通发货，全国大部分地区 48 小时内送达。")
doc.add_paragraph("新疆、西藏、青海等偏远地区需额外增加 3-5 天，且可能不支持货到付款。")
doc.add_paragraph("港澳台及海外地区暂不支持配送，如需寄送请联系客服协商。")
doc.save(DOCX_PATH)

with io.open(BAD_PATH, "w", encoding="utf-8") as f:
    f.write("a,b\n1,2\n")

client = httpx.Client(base_url=BASE, timeout=300.0)
report = {}


def stats_chunks():
    return client.get("/api/v1/stats").json().get("total_chunks")


def upload(path, label):
    name = os.path.basename(path)
    with io.open(path, "rb") as fh:
        files = {"file": (name, fh, "application/octet-stream")}
        t0 = time.perf_counter()
        r = client.post("/api/v1/ingest", files=files)
        dt = round(time.perf_counter() - t0, 2)
    try:
        body = r.json()
    except Exception:
        body = r.text[:200]
    return {"file": name, "status": r.status_code, "elapsed_s": dt, "body": body}


report["chunks_before"] = stats_chunks()
for path in (MD_PATH, TXT_PATH, DOCX_PATH):
    report[os.path.basename(path)] = upload(path, os.path.basename(path))
    report["chunks_after_" + os.path.basename(path)] = stats_chunks()
report["unsupported_csv"] = upload(BAD_PATH, "csv")

# 检索验证：问一个只有新导入的 md 里才有的问题
t0 = time.perf_counter()
r = client.post("/api/v1/chat", json={"message": "积分怎么获得？签到有奖励吗？", "session_id": "kb-test-points"})
report["query_points"] = {"status": r.status_code, "elapsed_s": round(time.perf_counter() - t0, 2), "body": r.json()}

t0 = time.perf_counter()
r = client.post("/api/v1/chat", json={"message": "开票需要多久？", "session_id": "kb-test-invoice"})
report["query_invoice"] = {"status": r.status_code, "elapsed_s": round(time.perf_counter() - t0, 2), "body": r.json()}

t0 = time.perf_counter()
r = client.post("/api/v1/chat", json={"message": "港澳台能配送吗？", "session_id": "kb-test-ship"})
report["query_ship"] = {"status": r.status_code, "elapsed_s": round(time.perf_counter() - t0, 2), "body": r.json()}

report["stats_final"] = client.get("/api/v1/stats").json()

with io.open(os.path.join(WORK, "ingest_report.json"), "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=2)
print("done ->", os.path.join(WORK, "ingest_report.json"))
print("chunks:", report["chunks_before"], "->", report["stats_final"]["total_chunks"])
